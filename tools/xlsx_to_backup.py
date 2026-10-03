#!/usr/bin/env python3
"""Convert the original '2026 Budget.xlsx' Google-Sheet export into the app's backup JSON
(docs/SYNC.md). Import the result from Settings -> Data -> Import backup in either app.

Usage: python3 tools/xlsx_to_backup.py "2026 Budget.xlsx" budget-import.json
The output contains personal financial data: do NOT commit it (the repo is public).
"""
import datetime as dt
import json
import sys
import uuid

import openpyxl

NS = uuid.UUID("6f1d3c3e-5b7a-4d7e-9a43-2f0c1b0e7a11")


def uid(*parts):
    return str(uuid.uuid5(NS, "/".join(str(p) for p in parts)))


CATS = [  # name, kind, tracking, multiplier, sort, sheet row
    ("Paychecks", "income", "ledger", 1, 0, 3),
    ("Rent", "expense", "manual", 1, 10, 5),
    ("Subscriptions", "expense", "ledger", 1, 11, 6),
    ("Food", "expense", "ledger", 1, 12, 7),
    ("Fun", "expense", "ledger", 1, 13, 8),
    ("Gas", "expense", "ledger", 1, 14, 9),
    ("Misc", "expense", "ledger", 1, 15, 10),
    ("Car Ins", "expense", "manual", 1, 16, 11),
    ("Utilities", "expense", "manual", 1, 17, 12),
    ("Phone Bill", "expense", "manual", 1, 18, 13),
    ("Roth", "savings", "manual", 1, 20, 15),
    ("401k", "savings", "manual", 2, 21, 16),
    ("Taxable Brokerage", "savings", "manual", 1, 22, 17),
    ("HSA", "savings", "manual", 1, 23, 18),
]
CAT_ID = {c[0]: uid("cat", c[0]) for c in CATS}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]
# ledger tables: category -> (date col, item col, amount col, first row, last row)
LEDGERS = {
    "Paychecks": ("G", "H", "I", 4, 5),
    "Subscriptions": ("G", "H", "I", 9, 13),
    "Misc": ("G", "H", "I", 17, 29),
    "Food": ("K", "L", "M", 4, 22),
    "Fun": ("O", "P", "Q", 4, 29),
    "Gas": ("K", "L", "M", 26, 29),
}


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 6)
    return None


def clean(s):
    return " ".join(str(s).split()) if s is not None else ""


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)  # cached values
    wf = openpyxl.load_workbook(src)  # formulas (to know which actuals were typed)
    out = {"app": "budget", "version": 1,
           "exported_at": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
           "settings": {"net_income": 68000, "gross_income": 85000, "currency": "USD"},
           "categories": [], "months": [], "budgets": [], "transactions": [], "recurring_items": [],
           "accounts": [], "ledger_entries": []}
    for name, kind, tracking, mult, sort, _ in CATS:
        out["categories"].append({"id": CAT_ID[name], "name": name, "kind": kind, "tracking": tracking,
                                  "match_multiplier": mult, "sort_order": sort, "icon": None,
                                  "color": None, "archived": False})
    year = int(wb["Summary"]["B2"].value or 2026)

    for ws_name in wb.sheetnames:
        if ws_name not in MONTHS:
            continue
        m = MONTHS.index(ws_name) + 1
        ws, wsf = wb[ws_name], wf[ws_name]
        out["months"].append({"year": year, "month": m, "closed": bool(ws["C23"].value), "note": None})
        for name, kind, tracking, mult, sort, row in CATS:
            expected = num(ws[f"C{row}"].value)
            raw = wsf[f"D{row}"].value
            typed = raw is not None and not (isinstance(raw, str) and raw.startswith("="))
            actual = num(raw) if typed else None
            if expected is None and actual is None:
                continue
            out["budgets"].append({"year": year, "month": m, "category_id": CAT_ID[name],
                                   "expected": expected, "actual": actual})
        for cat, (dc, ic, ac, r0, r1) in LEDGERS.items():
            for r in range(r0, r1 + 1):
                amount = num(ws[f"{ac}{r}"].value)
                item = clean(ws[f"{ic}{r}"].value)
                if amount is None and not item:
                    continue
                d = ws[f"{dc}{r}"].value
                date = d.date().isoformat() if isinstance(d, dt.datetime) else None
                out["transactions"].append({"id": uid("txn", year, m, cat, r), "year": year, "month": m,
                                            "category_id": CAT_ID[cat], "date": date, "item": item,
                                            "amount": amount or 0, "note": None})

    # Recurring subscriptions = the template rows pre-filled in the latest month
    last = max(out["months"], key=lambda x: x["month"])["month"]
    subs = [t for t in out["transactions"] if t["month"] == last and t["category_id"] == CAT_ID["Subscriptions"]]
    for i, t in enumerate(subs):
        out["recurring_items"].append({"id": uid("rec", t["item"]), "category_id": CAT_ID["Subscriptions"],
                                       "item": t["item"], "amount": t["amount"], "day_of_month": 1,
                                       "active": True, "sort_order": i})

    s = wb["Summary"]
    groups = {"Regions": ("cash", True), "SoFi": ("cash", True), "Cash": ("cash", True),
              "Venmo": ("cash", True), "Credit": ("debt", True), "SoFi HYSA": ("cash", True),
              "Brokerage": ("investment", False), "Roth IRA": ("investment", False),
              "401k": ("investment", False), "HSA": ("investment", False), "Car": ("asset", False)}
    for i, r in enumerate(range(3, 14)):
        name = s[f"G{r}"].value
        if not name:
            continue
        grp, liquid = groups.get(name, ("cash", False))
        acct = {"id": uid("acct", name), "name": name, "account_group": grp, "liquid": liquid,
                "balance": num(s[f"H{r}"].value) or 0, "linked_category_id": None, "base_amount": 0,
                "sort_order": i, "archived": False}
        if name == "401k":
            acct.update(balance=0, linked_category_id=CAT_ID["401k"], base_amount=0)
        if name == "HSA":
            acct.update(balance=0, linked_category_id=CAT_ID["HSA"], base_amount=500)
        out["accounts"].append(acct)
    for i, r in enumerate(range(3, 14)):
        name = s[f"J{r}"].value
        if name and num(s[f"K{r}"].value) is not None:
            out["ledger_entries"].append({"id": uid("ledger", name), "name": clean(name),
                                          "amount": num(s[f"K{r}"].value), "note": None,
                                          "settled": False, "sort_order": i})

    with open(dst, "w") as f:
        json.dump(out, f, indent=1)
    print(f"wrote {dst}: " + ", ".join(f"{k}={len(v)}" for k, v in out.items() if isinstance(v, list)))


if __name__ == "__main__":
    main(*sys.argv[1:3])
